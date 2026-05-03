import asyncio
import json
import mimetypes
import os
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote
from uuid import uuid4

import aiohttp
from aiohttp import web


PROJECT_ROOT = Path(__file__).resolve().parent
RUNS_DIR = PROJECT_ROOT / "exports" / "worker-runs"
EXPORT_BUCKET = os.environ.get("SUPABASE_EXPORT_BUCKET", "lead-exports")


def _json_response(payload, status=200):
    return web.json_response(payload, status=status)


def _check_secret(request):
    expected = os.environ.get("WORKER_API_SECRET") or os.environ.get("N8N_WEBHOOK_SECRET")
    if not expected:
        return True
    supplied = request.headers.get("x-exportflow-secret", "")
    return supplied == expected


def _parse_job_config(value):
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        try:
            parsed = json.loads(stripped)
        except json.JSONDecodeError:
            return None
        return parsed if isinstance(parsed, dict) else None
    return None


def _callback_secret():
    return os.environ.get("N8N_WEBHOOK_SECRET") or os.environ.get("WORKER_API_SECRET")


def _callback_headers():
    headers = {"Content-Type": "application/json"}
    secret = _callback_secret()
    if secret:
        headers["x-exportflow-secret"] = secret
    return headers


def _count_rows(file_path):
    if file_path.suffix == ".csv":
        try:
            with file_path.open("r", encoding="utf-8", errors="ignore") as file_handle:
                line_count = sum(1 for _line in file_handle)
            return max(line_count - 1, 0)
        except OSError:
            return None
    if file_path.suffix == ".json":
        try:
            with file_path.open("r", encoding="utf-8") as file_handle:
                payload = json.load(file_handle)
            return len(payload) if isinstance(payload, list) else None
        except (OSError, json.JSONDecodeError):
            return None
    if file_path.suffix == ".xlsx":
        workbook = None
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(file_path, read_only=True, data_only=True)
            worksheet = workbook.active
            return max((worksheet.max_row or 1) - 1, 0)
        except Exception:
            return None
        finally:
            if workbook:
                workbook.close()
    return None


def _resolve_run_path(path_value):
    path = Path(path_value)
    if path.is_absolute():
        return path
    return PROJECT_ROOT / path


def _prepare_job_config(job_config, job_id, run_dir):
    normalized = dict(job_config)
    normalized["job_id"] = str(normalized.get("job_id") or job_id)

    delivery = dict(normalized.get("delivery", {}))
    if not delivery.get("output_dir"):
        delivery["output_dir"] = str(run_dir / "exports")
    normalized["delivery"] = delivery

    _resolve_run_path(delivery["output_dir"]).mkdir(parents=True, exist_ok=True)
    return normalized


def _normalize_status_event(event, job_id):
    normalized = dict(event)
    normalized["job_id"] = str(normalized.get("job_id") or job_id)
    if normalized.get("status") == "delivered":
        normalized["status"] = "exporting"
        normalized["message"] = "Scraper finished; uploading lead export files."
    return normalized


def _lead_export_files(job_config, job_id):
    delivery = job_config.get("delivery", {}) if isinstance(job_config, dict) else {}
    output_dir = delivery.get("output_dir")
    if not output_dir:
        return []

    export_dir = _resolve_run_path(output_dir)
    leads_files = list(export_dir.glob(f"{job_id}_leads.*"))
    audit_files = list(export_dir.glob(f"{job_id}_audit.*"))
    candidates = sorted(leads_files + audit_files)
    return [path for path in candidates if path.is_file()]


def _job_status_from_event(status):
    return {
        "starting": "running",
        "discovering": "running",
        "discovered": "running",
        "analyzing": "running",
        "enriching": "running",
        "enriched": "running",
        "scoring": "running",
        "exporting": "running",
        "terminal": "running",
        "delivered": "delivered",
        "failed": "failed",
    }.get(status, "running")


async def _post_callback(url, payload):
    if not url:
        return
    async with aiohttp.ClientSession() as session:
        async with session.post(url, headers=_callback_headers(), json=payload, timeout=60) as response:
            if response.status >= 400:
                text = await response.text()
                raise RuntimeError(f"Callback failed with HTTP {response.status}: {text[:500]}")


async def _publish_status_event(job_id, event, status_callback=None):
    normalized = _normalize_status_event(event, job_id)
    if status_callback:
        await _post_callback(status_callback, normalized)
        return

    status = str(normalized.get("status") or "running")
    message = str(normalized.get("message") or "Worker event received.")
    mapped_status = _job_status_from_event(status)
    await _add_job_event(job_id, status, message, normalized)
    
    # Don't update the overall job status for every single terminal log line
    if status != "terminal":
        await _update_job_status(job_id, mapped_status, message if mapped_status == "failed" else None)


async def _forward_status_events(status_output_path, job_id, payload, process):
    status_callback = payload.get("status_callback")
    position = 0

    while True:
        if status_output_path.exists():
            with status_output_path.open("r", encoding="utf-8") as file_handle:
                file_handle.seek(position)
                lines = file_handle.readlines()
                position = file_handle.tell()

            for line in lines:
                stripped = line.strip()
                if not stripped:
                    continue
                try:
                    event = json.loads(stripped)
                except json.JSONDecodeError:
                    print(f"[worker] ignored malformed status event for job {job_id}: {stripped[:200]}", flush=True)
                    continue
                try:
                    await _publish_status_event(job_id, event, status_callback)
                except Exception as exc:
                    print(f"[worker] status event publish failed for job {job_id}: {exc}", flush=True)

        if process.poll() is not None:
            if not status_output_path.exists() or position >= status_output_path.stat().st_size:
                return

        await asyncio.sleep(1)


async def _upload_to_supabase(file_path, storage_path):
    supabase_url = os.environ.get("SUPABASE_URL") or os.environ.get("NEXT_PUBLIC_SUPABASE_URL")
    service_role_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url or not service_role_key:
        raise RuntimeError("Missing SUPABASE_URL and/or SUPABASE_SERVICE_ROLE_KEY on the worker.")

    encoded_path = quote(storage_path, safe="/")
    upload_url = f"{supabase_url.rstrip('/')}/storage/v1/object/{EXPORT_BUCKET}/{encoded_path}"
    content_type = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
    headers = {
        "Authorization": f"Bearer {service_role_key}",
        "apikey": service_role_key,
        "Content-Type": content_type,
        "x-upsert": "true",
    }
    async with aiohttp.ClientSession() as session:
        async with session.post(upload_url, headers=headers, data=file_path.read_bytes(), timeout=120) as response:
            if response.status >= 400:
                text = await response.text()
                raise RuntimeError(f"Supabase upload failed with HTTP {response.status}: {text[:500]}")


async def _supabase_rest_request(method, path, payload=None):
    supabase_url = os.environ.get("SUPABASE_URL") or os.environ.get("NEXT_PUBLIC_SUPABASE_URL")
    service_role_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url or not service_role_key:
        raise RuntimeError("Missing SUPABASE_URL and/or SUPABASE_SERVICE_ROLE_KEY on the worker.")

    url = f"{supabase_url.rstrip('/')}{path}" if path.startswith("/") else f"{supabase_url.rstrip('/')}/{path}"
    headers = {
        "Authorization": f"Bearer {service_role_key}",
        "apikey": service_role_key,
        "Content-Type": "application/json",
        "Prefer": "return=minimal",
    }
    async with aiohttp.ClientSession() as session:
        async with session.request(method, url, headers=headers, json=payload, timeout=60) as response:
            if response.status >= 400:
                text = await response.text()
                raise RuntimeError(f"Supabase REST {method} {path} failed with HTTP {response.status}: {text[:500]}")


async def _add_job_event(job_id, status, message, metadata=None):
    await _supabase_rest_request(
        "POST",
        "/rest/v1/job_events",
        {
            "job_id": job_id,
            "status": status,
            "message": message,
            "metadata": metadata or {},
        },
    )


async def _update_job_status(job_id, status, error_message=None, clear_error=False):
    payload = {"status": status}
    if error_message is not None or clear_error:
        payload["error_message"] = error_message
    await _supabase_rest_request(
        "PATCH",
        f"/rest/v1/lead_jobs?id=eq.{quote(job_id, safe='')}",
        payload,
    )


async def _register_exports(job_id, exports):
    rows = [
        {
            "job_id": job_id,
            "format": export["format"],
            "storage_path": export["storage_path"],
            "row_count": export.get("row_count"),
        }
        for export in exports
    ]
    await _supabase_rest_request("DELETE", f"/rest/v1/lead_exports?job_id=eq.{quote(job_id, safe='')}")
    await _supabase_rest_request("POST", "/rest/v1/lead_exports", rows)
    await _update_job_status(job_id, "delivered", clear_error=True)
    await _add_job_event(job_id, "delivered", "Lead export files uploaded to Supabase Storage.", {"exports": rows})


async def _mark_failed(job_id, message, metadata=None):
    await _update_job_status(job_id, "failed", message)
    await _add_job_event(job_id, "failed", message, metadata or {})


async def _deliver_exports(job_id, job_config, payload):
    files = _lead_export_files(job_config, job_id)
    if not files:
        raise RuntimeError("No lead export files were found to upload.")

    exports = []
    for file_path in files:
        storage_path = f"{job_id}/{file_path.name}"
        await _upload_to_supabase(file_path, storage_path)
        exports.append(
            {
                "format": file_path.suffix.lstrip(".") or "file",
                "storage_path": storage_path,
                "row_count": _count_rows(file_path),
            }
        )

    export_callback = payload.get("export_callback")
    if export_callback:
        await _post_callback(export_callback, {"job_id": job_id, "exports": exports})
        return

    await _register_exports(job_id, exports)


async def _monitor_process(process, job_id, job_config, payload, status_output_path, stdout_handle, stderr_handle, stdout_path):
    status_callback = payload.get("status_callback")
    status_forwarder = asyncio.create_task(_forward_status_events(status_output_path, job_id, payload, process))
    try:
        return_code = await asyncio.to_thread(process.wait)
    finally:
        stdout_handle.close()
        if stderr_handle:
            stderr_handle.close()
        try:
            await asyncio.wait_for(status_forwarder, timeout=10)
        except Exception as exc:
            print(f"[worker] status forwarder did not finish cleanly for job {job_id}: {exc}", flush=True)

    if return_code != 0:
        message = f"Worker process exited with code {return_code}."
        if status_callback:
            try:
                await _post_callback(status_callback, {"status": "failed", "message": message, "job_id": job_id})
                return
            except Exception as exc:
                print(f"[worker] failed callback failed for job {job_id}: {exc}", flush=True)
        try:
            await _mark_failed(job_id, message)
        except Exception as exc:
            print(f"[worker] failed status update failed for job {job_id}: {exc}", flush=True)
        return

    try:
        await _deliver_exports(job_id, job_config, payload)
        print(f"[worker] delivered exports for job {job_id}", flush=True)
    except Exception as exc:
        print(f"[worker] delivery failed for job {job_id}: {exc}", flush=True)
        message = "Worker finished but export delivery failed."
        metadata = {"job_id": job_id, "error": str(exc)}
        if status_callback:
            try:
                await _post_callback(status_callback, {"status": "failed", "message": message, **metadata})
                return
            except Exception as callback_exc:
                print(f"[worker] delivery failure callback failed for job {job_id}: {callback_exc}", flush=True)
        try:
            await _mark_failed(job_id, message, metadata)
        except Exception as mark_exc:
            print(f"[worker] delivery failure status update failed for job {job_id}: {mark_exc}", flush=True)


async def health(_request):
    return _json_response({"ok": True, "service": "exportflow-worker"})


async def stream_logs(request):
    job_id = request.match_info.get("job_id")
    log_path = RUNS_DIR / job_id / "stdout.log"

    if not log_path.exists():
        return web.Response(text="Log file not found", status=404)

    response = web.StreamResponse(
        status=200,
        reason='OK',
        headers={
            'Content-Type': 'text/event-stream',
            'Cache-Control': 'no-cache',
            'Connection': 'keep-alive',
            'Access-Control-Allow-Origin': '*',
        },
    )
    await response.prepare(request)

    try:
        with log_path.open("r", encoding="utf-8", errors="replace") as file_handle:
            # Send current content
            while True:
                line = file_handle.readline()
                if not line:
                    break
                event_data = json.dumps({"message": line.strip()})
                await response.write(f"data: {event_data}\n\n".encode('utf-8'))

            # Tail for new content
            while True:
                line = file_handle.readline()
                if not line:
                    # Check if process is done? For now just sleep
                    await asyncio.sleep(0.5)
                    continue
                event_data = json.dumps({"message": line.strip()})
                await response.write(f"data: {event_data}\n\n".encode('utf-8'))
    except ConnectionResetError:
        pass
    except Exception as exc:
        print(f"[worker] log stream error for job {job_id}: {exc}")

    return response


async def run_job(request):
    print(f"[worker] /run-job hit at {datetime.now(timezone.utc).isoformat()}", flush=True)
    if not _check_secret(request):
        print("[worker] rejected request: invalid secret", flush=True)
        return _json_response({"error": "Invalid worker secret."}, status=401)

    try:
        payload = await request.json()
    except json.JSONDecodeError:
        print("[worker] rejected request: body was not JSON", flush=True)
        return _json_response({"error": "Request body must be JSON."}, status=400)

    job_config = _parse_job_config(payload.get("job_config")) or _parse_job_config(payload)
    job_id = str(payload.get("job_id") or (job_config or {}).get("job_id") or uuid4()).strip()
    if not job_config:
        print(f"[worker] rejected request: missing job_config | keys={list(payload.keys())}", flush=True)
        return _json_response({"error": "job_id and job_config are required."}, status=400)

    print(f"[worker] accepted job {job_id}", flush=True)
    run_dir = RUNS_DIR / job_id
    run_dir.mkdir(parents=True, exist_ok=True)
    job_config = _prepare_job_config(job_config, job_id, run_dir)

    job_config_path = run_dir / "job_config.json"
    status_output_path = run_dir / "events.jsonl"
    stdout_path = run_dir / "stdout.log"
    stderr_path = run_dir / "stderr.log"

    with job_config_path.open("w", encoding="utf-8") as file_handle:
        json.dump(job_config, file_handle, ensure_ascii=True, indent=2)

    command = [
        sys.executable,
        "-u",
        str(PROJECT_ROOT / "final scrapper.py"),
        "--job-config",
        str(job_config_path),
        "--status-output",
        str(status_output_path),
        "--job-id",
        job_id,
    ]

    stdout_handle = stdout_path.open("ab")
    process = subprocess.Popen(
        command,
        cwd=str(PROJECT_ROOT),
        stdout=stdout_handle,
        stderr=subprocess.STDOUT,  # Merge stderr into stdout for live terminal logs
        start_new_session=os.name != "nt",
    )
    asyncio.create_task(_monitor_process(process, job_id, job_config, payload, status_output_path, stdout_handle, None, stdout_path))

    started_at = datetime.now(timezone.utc).isoformat()
    return _json_response(
        {
            "ok": True,
            "job_id": job_id,
            "pid": process.pid,
            "started_at": started_at,
            "run_dir": str(run_dir),
            "status_output": str(status_output_path),
            "stdout": str(stdout_path),
            "stderr": str(stderr_path),
        },
        status=202,
    )


def create_app():
    app = web.Application()
    app.router.add_get("/health", health)
    app.router.add_get("/api/logs/{job_id}", stream_logs)
    app.router.add_post("/run-job", run_job)
    return app


if __name__ == "__main__":
    port = int(os.environ.get("WORKER_PORT", "8787"))
    web.run_app(create_app(), host="0.0.0.0", port=port)
