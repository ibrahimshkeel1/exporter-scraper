import csv
import json
import os
from datetime import datetime, timezone

class LeadExport:
    def __init__(self, output_file="leads.csv"):
        self.output_file = output_file
        self.headers = [
            "run_id",
            "scraped_at",
            "region",
            "industry",
            "company_name",
            "domain",
            "website",
            "discovery_url",
            "source_name",
            "source_url",
            "discovery_method",
            "analyzed_index",
            "elapsed_seconds",
            "is_a_plus",
            "qualified",
            "passes_hard_checks",
            "export_eligible",
            "buyer_type",
            "lead_pack_status",
            "manual_review_required",
            "product_fit",
            "product_evidence",
            "buyer_evidence",
            "buyer_side_evidence",
            "sales_side_evidence",
            "contact_evidence",
            "contact_route",
            "outreach_contact",
            "evidence_url",
            "negative_evidence",
            "disqualification_reasons",
            "recommended_pitch_angle",
            "lead_summary",
            "closeability_notes",
            "linkedin_url",
            "social_urls",
            "contact_form_urls",
            "emails",
            "high_quality_emails",
            "email_quality",
            "rejected_emails",
            "fetch_ok",
            "fetch_status_codes",
            "fetch_errors",
            "crawled_pages",
            "score",
            "tier",
            "qualification_reasons",
            "score_breakdown",
        ]

    @staticmethod
    def _output_path(base_output_file, extension):
        stem, _ = os.path.splitext(base_output_file)
        if not stem:
            stem = base_output_file
        return f"{stem}{extension}"

    @staticmethod
    def _ensure_parent_dir(file_path):
        parent_dir = os.path.dirname(file_path)
        if parent_dir:
            os.makedirs(parent_dir, exist_ok=True)

    @staticmethod
    def _company_name_from_domain(domain):
        if not domain:
            return "Unknown"
        label = domain.split(".")[0].replace("-", " ")
        return " ".join(piece.capitalize() for piece in label.split())

    def _format_lead(self, lead, region, industry, run_id):
        return {
            "run_id": run_id,
            "scraped_at": datetime.now(timezone.utc).isoformat(),
            "region": region,
            "industry": industry,
            "company_name": self._company_name_from_domain(lead.get("domain", "")),
            "domain": lead.get("domain", ""),
            "website": lead.get("url", ""),
            "discovery_url": lead.get("discovery_url", ""),
            "source_name": lead.get("source_name", ""),
            "source_url": lead.get("source_url", ""),
            "discovery_method": lead.get("discovery_method", ""),
            "analyzed_index": lead.get("analyzed_index", ""),
            "elapsed_seconds": lead.get("elapsed_seconds", ""),
            "is_a_plus": lead.get("is_a_plus", False),
            "qualified": lead.get("qualified", False),
            "passes_hard_checks": lead.get("passes_hard_checks", False),
            "export_eligible": lead.get("export_eligible", False),
            "buyer_type": lead.get("buyer_type", ""),
            "lead_pack_status": lead.get("lead_pack_status", ""),
            "manual_review_required": lead.get("manual_review_required", True),
            "product_fit": lead.get("product_fit", ""),
            "product_evidence": lead.get("product_evidence", ""),
            "buyer_evidence": lead.get("buyer_evidence", ""),
            "buyer_side_evidence": lead.get("buyer_side_evidence", ""),
            "sales_side_evidence": lead.get("sales_side_evidence", ""),
            "contact_evidence": lead.get("contact_evidence", ""),
            "contact_route": lead.get("contact_route", ""),
            "outreach_contact": lead.get("outreach_contact", ""),
            "evidence_url": lead.get("evidence_url", ""),
            "negative_evidence": lead.get("negative_evidence", ""),
            "disqualification_reasons": lead.get("disqualification_reasons", ""),
            "recommended_pitch_angle": lead.get("recommended_pitch_angle", ""),
            "lead_summary": lead.get("lead_summary", ""),
            "closeability_notes": lead.get("closeability_notes", ""),
            "linkedin_url": lead.get("linkedin_url", ""),
            "social_urls": ", ".join(lead.get("social_urls", [])),
            "contact_form_urls": " | ".join(lead.get("contact_form_urls", [])),
            "emails": ", ".join(lead.get("emails", [])),
            "high_quality_emails": ", ".join(lead.get("high_quality_emails", [])),
            "email_quality": lead.get("email_quality", "none"),
            "rejected_emails": ", ".join(lead.get("rejected_emails", [])),
            "fetch_ok": lead.get("fetch_ok", False),
            "fetch_status_codes": ",".join(str(code) for code in lead.get("fetch_status_codes", [])),
            "fetch_errors": " | ".join(lead.get("fetch_errors", [])),
            "crawled_pages": " | ".join(lead.get("crawled_pages", [])),
            "score": lead.get("score", 0),
            "tier": lead.get("tier", "C"),
            "qualification_reasons": lead.get("qualification_reasons", ""),
            "score_breakdown": json.dumps(lead.get("score_breakdown", {}), ensure_ascii=True),
        }

    def save_to_csv(self, leads, region, industry, run_id):
        if not leads:
            print("No leads to save.")
            return

        formatted_leads = [
            self._format_lead(lead, region=region, industry=industry, run_id=run_id)
            for lead in leads
        ]

        csv_file = self._output_path(self.output_file, ".csv")
        self._ensure_parent_dir(csv_file)
        with open(csv_file, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=self.headers)
            writer.writeheader()
            writer.writerows(formatted_leads)

        print(f"Saved {len(leads)} leads to {csv_file}")

    def save_to_json(self, leads, region, industry, run_id):
        if not leads:
            print("No leads to save.")
            return
        formatted_leads = [
            self._format_lead(lead, region=region, industry=industry, run_id=run_id)
            for lead in leads
        ]
        json_file = self._output_path(self.output_file, ".json")
        self._ensure_parent_dir(json_file)
        with open(json_file, "w", encoding="utf-8") as file_handle:
            json.dump(formatted_leads, file_handle, ensure_ascii=True, indent=2)
        print(f"Saved {len(leads)} leads to {json_file}")

    def save_to_xlsx(self, leads, region, industry, run_id):
        if not leads:
            print("No leads to save.")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "Missing dependency: openpyxl. Install with `pip install -r scraper/requirements.txt`."
            ) from exc

        formatted_leads = [
            self._format_lead(lead, region=region, industry=industry, run_id=run_id)
            for lead in leads
        ]

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Verified Leads"
        worksheet.append(self.headers)

        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        for lead in formatted_leads:
            worksheet.append([lead.get(header, "") for header in self.headers])

        worksheet.freeze_panes = "A2"
        for index, header in enumerate(self.headers, start=1):
            max_length = len(header)
            for row in worksheet.iter_rows(min_col=index, max_col=index, min_row=2):
                value = row[0].value
                if value is not None:
                    max_length = max(max_length, min(len(str(value)), 60))
            worksheet.column_dimensions[get_column_letter(index)].width = min(max_length + 2, 64)

        xlsx_file = self._output_path(self.output_file, ".xlsx")
        self._ensure_parent_dir(xlsx_file)
        workbook.save(xlsx_file)
        print(f"Saved {len(leads)} leads to {xlsx_file}")

    def save(self, leads, region, industry, run_id, output_format="csv"):
        if output_format == "csv":
            self.save_to_csv(leads, region, industry, run_id)
        elif output_format == "json":
            self.save_to_json(leads, region, industry, run_id)
        elif output_format == "xlsx":
            self.save_to_xlsx(leads, region, industry, run_id)
        elif output_format == "all":
            self.save_to_csv(leads, region, industry, run_id)
            self.save_to_json(leads, region, industry, run_id)
            self.save_to_xlsx(leads, region, industry, run_id)
        else:
            self.save_to_csv(leads, region, industry, run_id)
            self.save_to_json(leads, region, industry, run_id)
